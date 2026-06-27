import io
from datetime import datetime, timedelta
from flask import render_template


def _weasy_pdf(html):
    try:
        from weasyprint import HTML
        return HTML(string=html).write_pdf()
    except Exception:
        return ("PDF engine unavailable. Install WeasyPrint system dependencies.\n\n"
                + _strip_html(html)).encode("utf-8")


def _strip_html(html):
    import re
    text = re.sub(r"<[^>]+>", " ", html)
    return re.sub(r"\s+", " ", text).strip()


def generate_most_pdf(study_id):
    from app.models.most_study import MostStudy
    study = MostStudy.query.get(study_id)
    html = render_template("reports_pdf/most_study.html", study=study,
                           now=datetime.utcnow())
    return _weasy_pdf(html)


def generate_capa_pdf(capa_id):
    from app.models.capa import CapaRecord
    capa = CapaRecord.query.get(capa_id)
    html = render_template("reports_pdf/capa.html", capa=capa, now=datetime.utcnow())
    return _weasy_pdf(html)


def generate_rejections_excel(plant_id, start=None, end=None):
    from openpyxl import Workbook
    from app.models.rejection import RejectionEvent
    from app.models.shift import Shift
    wb = Workbook()
    ws = wb.active
    ws.title = "Rejections"
    ws.append(["ID", "Shift", "Defect Type", "Qty Rejected", "Probable Cause",
               "Check Point", "Batch", "Logged At"])
    q = RejectionEvent.query.join(Shift).join(Shift.line).filter_by(plant_id=plant_id)
    if start:
        q = q.filter(RejectionEvent.created_at >= start)
    if end:
        q = q.filter(RejectionEvent.created_at <= end)
    for r in q.order_by(RejectionEvent.created_at.desc()).all():
        ws.append([r.id, r.shift.shift_label if r.shift else "", r.defect_type,
                   r.quantity_rejected, r.probable_cause or "", r.quality_check_point or "",
                   r.batch_reference or "", r.created_at.strftime("%Y-%m-%d %H:%M")])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_downtime_excel(plant_id, start=None, end=None):
    from openpyxl import Workbook
    from app.models.downtime import DowntimeEvent
    from app.models.machine import Machine
    wb = Workbook()
    ws = wb.active
    ws.title = "Downtime"
    ws.append(["ID", "Machine", "Event Type", "Reason", "Start", "End",
               "Duration (min)", "Resolved"])
    mids = [m.id for m in Machine.query.join(Machine.line).filter_by(plant_id=plant_id).all()]
    q = DowntimeEvent.query.filter(DowntimeEvent.machine_id.in_(mids))
    if start:
        q = q.filter(DowntimeEvent.start_time >= start)
    if end:
        q = q.filter(DowntimeEvent.start_time <= end)
    for d in q.order_by(DowntimeEvent.start_time.desc()).all():
        ws.append([d.id, d.machine.machine_name if d.machine else "", d.event_type,
                   d.reason_code.label if d.reason_code else "",
                   d.start_time.strftime("%Y-%m-%d %H:%M"),
                   d.end_time.strftime("%Y-%m-%d %H:%M") if d.end_time else "",
                   round(d.duration_minutes, 1) if d.duration_minutes else "",
                   "Yes" if d.is_resolved else "No"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def generate_report(report_type, params, format="pdf"):
    plant_id = params.get("plant_id")
    start = params.get("start_date")
    end = params.get("end_date")
    if report_type == "rejection_trend":
        return generate_rejections_excel(plant_id, start, end), "xlsx", "rejection_trend"
    if report_type == "downtime":
        return generate_downtime_excel(plant_id, start, end), "xlsx", "downtime"
    html = render_template("reports_pdf/generic.html", report_type=report_type,
                           params=params, now=datetime.utcnow())
    if format == "excel":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.append([report_type, str(start), str(end)])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf.read(), "xlsx", report_type
    return _weasy_pdf(html), "pdf", report_type
